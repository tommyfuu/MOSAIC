"""Write Supplemental Table 4 (PubTator key-term counts per significant taxon) as an Excel workbook,
in the layout of SuppTable4.pubtator_merged_outputs_110624.xlsx, with live Excel formulas.

Needs openpyxl (not in the pipeline conda envs: pip install openpyxl).
Usage: python3 make_supp_table4.py [--succinct]   (--succinct drops bat/ruv3/svd/rbe/bmc)
Inputs: trial_rw/outputs/{dataset}_taxa_union_wilcoxon.csv and trial_rw/outputs/pubtator_keyterms/{dataset}__{term}.csv
"""
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as col

sys.path.insert(1, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from paths import REPO_ROOT

IN = os.path.join(REPO_ROOT, 'trial_rw', 'outputs')
SUCCINCT = '--succinct' in sys.argv  # drop the methods left out of the succinct plots
SUCCINCT_DROP = {'bat', 'ruv3', 'svd', 'rbe', 'bmc'}
OUT = os.path.join(IN, 'pubtator_keyterms', 'SuppTable4_pubtator_keyterms' + ('_succinct' if SUCCINCT else '') + '.xlsx')
TEN = '>=10'  # Methods text: "at least ten" publications (the Nov 2024 sheet used ">10")

COUNT_METHODS = ['nobc', 'harmony', 'combat_seq', 'limma', 'MMUPHin', 'ConQuR', 'ConQuR_libsize', 'percentile_norm',
                 'debiasm', 'ruv3', 'bmc', 'bat', 'rbe', 'svd']
RELAB_METHODS = ['nobc', 'harmony', 'combat', 'limma', 'MMUPHin', 'ConQuR_rel', 'percentile_norm',
                 'debiasm', 'ruv3', 'bmc', 'bat', 'rbe', 'svd']
SHEETS = [
    ('autism', 'autism_2_microbiomeHD', COUNT_METHODS,
     ['autism', 'autism spectrum disorders', 'gut-brain axis', 'neurological disorders', 'nervous systems']),
    ('cdi', 'cdi_3_microbiomeHD_clean', COUNT_METHODS,
     ['Clostridioides difficile', 'diarrhea', 'antibiotics', 'toxins', 'infection']),
    ('ibd', 'ibd_3_CMD', RELAB_METHODS,
     ['inflammatory bowel disease', "Crohn's disease", 'ulcerative colitis', 'abdominal pain', 'intestinal inflammation']),
    ('crc', 'crc_11_CMD', RELAB_METHODS,
     ['colon cancer', 'colorectal cancer', 'tumorigenesis', 'rectal bleeding', 'MSI subtype']),
]


def slug(term):
    return ''.join(ch for ch in term.replace(' ', '_') if ch.isalnum() or ch in '_-')


wb = Workbook()
main = wb.active
main.title = 'Main'
for line in [
    'Supplemental Table 4. Number of publications returned when searching each differentially abundant taxon along with a pre-defined clinical key term.',
    'For each real-world disease data, all significant taxa identified via differential abundance test in at least one of the pre-integration (nobc) or post-integration datasets are included to be searched with the pre-defined terms. The dataset (strategy) a taxon comes from is noted to the left of each subsheet.',
    'For each taxon, we also calculate the number of clinical terms that return at least 1 or at least 10 publications when searched with the taxon. Taxa that have at least 5 such 1-publication terms or at least 3 such 10-publication terms are "strongly-relevant taxa", as defined in Methods.',
    'To use this approach to evaluate integration strategies, we also calculate the total number of strongly-relevant taxa identified post each integration strategy as a proxy to examine the ability of each strategy in conserving true biological signals.',
]:
    main.append([line])

for sheet, ds, methods, terms in SHEETS:
    union = pd.read_csv(f'{IN}/{ds}_taxa_union_wilcoxon.csv', index_col=0)
    methods = [m for m in methods if m in union.columns and not (SUCCINCT and m in SUCCINCT_DROP)]
    union = union[union[methods].astype(bool).any(axis=1)]  # taxa significant in at least one listed method
    counts = {}
    for t in terms:
        c = pd.read_csv(f'{IN}/pubtator_keyterms/{ds}__{slug(t)}.csv')
        counts[t] = dict(zip(c['search_term'], c['pubtator_count']))
    ws = wb.create_sheet(sheet)
    m0, nm = 3, len(methods)                      # method booleans start at column C
    c_gt0, c_ge10 = m0 + nm, m0 + nm + 1
    t0 = m0 + nm + 2                              # 5 term columns
    c_sr = t0 + 5
    k0 = c_sr + 1                                 # per-method strongly-relevant indicators
    header = [None, 'cleaned_taxa'] + methods + ['#terms with >0 hit', f'#terms with {TEN} hit'] + terms + \
             ['strongly relevant'] + [f'{m}_count' for m in methods]
    ws.append(header)
    for i, (raw, row) in enumerate(union.iterrows(), start=2):
        tr = f'{col(t0)}{i}:{col(t0 + 4)}{i}'
        K, L = f'{col(c_gt0)}{i}', f'{col(c_ge10)}{i}'
        vals = [raw, row['cleaned_taxa']] + [bool(row[m]) for m in methods] + \
               [f'=COUNTIF({tr}, ">0")', f'=COUNTIF({tr}, "{TEN}")'] + \
               [counts[t].get(row['cleaned_taxa']) for t in terms] + \
               [f'=IF(OR({K}=5, {L}>=3), 1, 0)'] + \
               [f'=IF(AND({col(m0 + j)}{i}=TRUE, OR({K}>=5, {L}>=3)), 1, 0)' for j in range(nm)]
        ws.append(vals)
    last = len(union) + 1
    ws.append([None] * (k0 - 1) + [f'=SUM({col(k0 + j)}2:{col(k0 + j)}{last})' for j in range(nm)])
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 32
    ws.freeze_panes = 'C2'

wb.save(OUT)
print(OUT)
